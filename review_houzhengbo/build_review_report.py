from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "HouZhengbo_P1_论文审阅问题清单.docx"

BLUE = "2E5E8C"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
TEXT = "202124"
MUTED = "5E6670"
RED = "9B1C1C"


CRITICAL = [
    ("P9", "LIST OF TABLES", "表题与正文不一致，页码也错误。清单写的是 “Comparison of the RAG System with Other AI Models, 6”，正文表题是 “Comparison of Large Language Models and RAG-based Systems for University Student Information Services”，实际正文页为 13（PDF P24）。", "更新表目录字段，确保标题和页码与 Table 2.1 完全一致。"),
    ("P10", "Pictures No. / Page", "插图目录末尾出现空的 “Pictures No. Page” 栏目。", "若没有 Pictures，删除这两行。"),
    ("P26", "/", "Figure 2.1 上方有一个孤立的斜杠。", "删除斜杠。"),
    ("P33-P36", "Figure 3.3 references", "3.5 节多处把 RAG 工作流引用为 Figure 3.3，但 Figure 3.3 实际是跨语言查询对齐图；RAG 工作流在 Figure 3.4(a)/(b)。", "Agent/tool decision 应指 Figure 3.4(a)；hybrid retrieval 和 answer generation 应指 Figure 3.4(b)。"),
    ("P17/P41", "Chapter 4 and Chapter 5", "正文和目录目前只有 Chapter 1-3，但 1.8 写 “five chapters”，3.8 又引用 Chapter 4。", "若这是 P1/Proposal，明确使用将来时并说明后续章节；若是当前提交稿，目录与章节说明必须一致。"),
    ("P27-P41", "Methodology tense", "方法章节交替使用 will include、is implemented、were implemented、provides、started 等时态。", "已完成系统用一般过去时/现在时统一；尚未完成的 P1 方案用一般将来时统一。不要混用。"),
    ("P18-P41", "Citation style", "正文括号引用混用 “&” 和 “and”；部分句号放在引用前，如 “models. (Lewis et al., 2020).”", "按学院指定格式统一。若采用 Harvard/UKM 常见格式，正文叙述用 and，括号内也按模板统一；引用应放在句号前。"),
    ("P13/P14/P27/P36/P41", "Absolute claims", "多处声称 RAG “without introducing any hallucinations / prevents hallucinations / without any hallucinations / factually correct”。", "改为 reduces the risk of hallucination、improves grounding 或 supports factual verification，除非有实验能证明绝对结论。"),
    ("P11", "List of abbreviations", "缩略语表缺少正文频繁使用的 BM25、OCR、RRF、SSE、TTL、JSON、YAML、SDK、UUID 等。FTSM 的马来文全称还使用了 “&”，与封面 “dan” 不一致。", "补齐重要缩略语；改为 “Fakulti Teknologi dan Sains Maklumat”。"),
    ("P18-P45", "Terminology/case", "RAGs / RAG system、Langchain / LangChain、Qwen / Tongyi Qwen / Qwen3-Max、backend / back-end、Section / section 等不统一。", "建立术语表并全文统一：RAG systems、LangChain、Qwen3-Max（具体模型）、backend、Section 3.x。"),
    ("P18/P24/P29/P33", "Heading/caption styles", "2.5 标题不是全大写；Figure 3.1 caption 使用普通正文样式，而其他图题使用 Caption；Table 2.1 的目录标题与正文题名不一致。", "统一标题大小写和 Caption 样式，并更新所有目录字段。"),
    ("P42-P45", "Reference formatting", "参考文献全部使用 Normal 样式，无悬挂缩进；页码范围混用 hyphen/en dash；正文引用格式也不统一。", "按 UKM 模板统一参考文献样式、悬挂缩进、行距、页码连接号和作者连接方式。"),
]


GLOBAL = [
    ("全文", "Articles", "多处漏冠词，例如 “designing question answering system”, “using DashScope model”, “via POST /api/chat endpoint”。", "逐句补充 a/an/the：designing a question-answering system；using the DashScope model；via the POST ... endpoint。"),
    ("全文", "allow/enable", "多处使用 “allow retrieving / allows increasing / enable comparing”，表达生硬。", "改为 enable retrieval、help increase、enable comparison，或 allow + object + to do。"),
    ("全文", "Nominalization", "“the process of generation of an answer”“the process of processing multilingual queries”等名词化过重。", "改为 answer generation、multilingual query processing。"),
    ("全文", "RAG plural", "把技术本身写成 “RAGs”。", "指系统时写 RAG systems；指方法时写 RAG。"),
    ("全文", "Programme/program", "学术课程与软件程序的 program 用法未区分。", "马来西亚英式英语中，学术课程用 programme；软件程序用 program。"),
    ("全文", "Question answering", "question answering 作定语时有时无连字符。", "统一为 question-answering system；单独作名词时写 question answering。"),
    ("全文", "Code identifiers", "文件名、函数、环境变量、接口路径和配置键与普通正文混排。", "统一使用等宽字体或 Word 的代码样式，例如 web_app.py、CHAT_MODEL_NAME、/api/chat。"),
    ("全文", "Section/Figure/Table", "章节、图表交叉引用的大小写不统一。", "引用具体编号时统一写 Section 3.4、Figure 3.3、Table 2.1。"),
    ("全文", "First/Second", "混用 First/Firstly、Second/Secondly。", "选一种并全文统一；正式论文建议 First, Second, Finally。"),
    ("全文", "Informal phrasing", "出现 “talks about”“As seen”“effects ... are obvious”“one needs to mention”等口语或主观表达。", "改为 describes、These results indicate、The component reduces...、The study covers...。"),
    ("全文", "Long sentences", "多个句子超过 45-60 词，并串联 3-5 个动作，主语不清。", "拆成两到三句，每句只保留一个主要动作或结论。"),
    ("全文", "British/American consistency", "存在 analyse/analyze、program/programme、pre-trained/pretrained 等潜在风格混用。", "以 UKM 模板要求为准统一拼写；文献原题保持原样。"),
]


FRONT = [
    ("P6", "has made more and more foreign students studying", "结构错误。", "has led to an increasing number of international students studying"),
    ("P6", "there are many Chinese ones", "ones 指人不自然。", "many are Chinese international students"),
    ("P6", "the new culture, institutions' rules", "institutions' rules 生硬，且 new culture 指代不清。", "an unfamiliar culture, institutional regulations"),
    ("P6", "FTSM of UKM University", "UKM 已含 University，重复且机构关系表达错误。", "FTSM at Universiti Kebangsaan Malaysia (UKM)"),
    ("P6", "without having any support to help Chinese international students", "冗长且搭配不自然。", "with limited language support for Chinese international students"),
    ("P6", "integrate the documents ... and experiences ... to create the database", "并列对象和目的表达不顺。", "integrate authoritative documents and students' experiential knowledge into a domain-specific knowledge base"),
    ("P6", "Chinese International students", "普通名词不应大写 International。", "Chinese international students"),
    ("P6", "the technology ... is supposed to be used", "语气不确定、被动冗长。", "the system uses semantic retrieval and large language models"),
    ("P6", "Besides, the ability to have multi-turn interactions ... is also provided", "Besides 偏口语，句子被动。", "The system also supports natural-language, multi-turn interaction."),
    ("P6", "as the result of the case study, one will be able to find", "冠词错误且 one 生硬。", "As a result, users will be able to retrieve"),
    ("P6", "academic schedule, supervisors' contacts, the processes of the application procedure", "名词数和搭配不自然。", "academic schedules, supervisors' contact details, application procedures"),
    ("P6", "the problem of hallucinations is solved", "绝对化结论，且尚无结果章节支撑。", "the risk of hallucination is reduced"),
    ("P11", "FTSM ... Fakulti Teknologi & Sains Maklumat", "与封面马来文名称不一致。", "Fakulti Teknologi dan Sains Maklumat"),
]


CHAPTER_1 = [
    ("P12", "a surge of relevance of the intelligent question answering systems", "搭配错误且冠词过多。", "a surge in the relevance of intelligent question-answering systems"),
    ("P12", "improve the information access and communication within the different domains", "information access 通常不加 the；different 冗余。", "improve information access and communication across domains"),
    ("P12", "in particular domain-based environment", "缺冠词，domain-based environment 生硬。", "in a specific domain"),
    ("P12", "students need to get accurate information", "get 偏口语。", "students need to obtain accurate information"),
    ("P12", "manual searching of information from numerous...", "名词化且搭配不自然。", "manually searching numerous web pages, notices, and documents"),
    ("P13", "suggests designing and implementing the ... system designed for", "designed 重复；research suggests 语气也弱。", "proposes the design and implementation of a ... system for"),
    ("P13", "integration of external knowledge into RAG system", "缺冠词。", "the integration of external knowledge into a RAG system"),
    ("P13", "are included into the architecture", "搭配错误。", "are included in the architecture"),
    ("P13", "this research will be aimed at designing", "冗长。", "this research aims to design"),
    ("P13", "Technology , Universiti", "逗号前有多余空格。", "Technology, Universiti"),
    ("P13", "has resulted in the gradual increase in international student populations", "冠词/数表达不自然。", "has resulted in a gradual increase in the international student population"),
    ("P13", "one of the major international groups of students", "重复 students。", "one of the largest international student groups"),
    ("P13", "allocation of supervisors", "搭配不自然。", "supervisor assignment"),
    ("P13", "these data are not formalized and structured", "experiential knowledge 通常不可数。", "this knowledge is not formally documented or structured"),
    ("P13", "large language models' abilities", "表达生硬。", "the capabilities of large language models"),
    ("P13", "RAG will allow retrieving", "allow 用法错误。", "RAG will enable the retrieval of / will allow the system to retrieve"),
    ("P14", "have not been fine-tuned based on any knowledge source of institutional documents", "搭配错误。", "have not been fine-tuned on institutional documents or other faculty-specific sources"),
    ("P14", "Faculty guidelines", "非正式专名不应大写。", "faculty guidelines"),
    ("P14", "high precision and factual accuracy is a must", "主谓不一致且偏口语。", "high precision and factual accuracy are essential"),
    ("P14", "previous research shows that work ... has primarily paid attention to", "research/work 重复且搭配错误。", "previous studies have primarily focused on"),
    ("P14", "students at universities who are natives of China", "冗长且不自然。", "Chinese international university students"),
    ("P14", "Especially, there is little research", "句首连接词不自然。", "In particular, little research"),
    ("P14", "the problem addressed in the current paper is related to", "冗长，paper 与 thesis 不一致。", "This study addresses"),
    ("P15", "from the student side", "中式表达。", "for students / from students' perspective"),
    ("P15", "relevance of answer retrieval in the multilingual setting of the university question answering", "语义不清。", "relevance of retrieved evidence and generated answers in a multilingual university setting"),
    ("P15", "speed of retrieving information and the efficiency of using the API", "指标表达不够明确。", "retrieval latency and API usage efficiency"),
    ("P15", "The goal... / The objective... / The aim...", "三个目标结构不平行。", "统一为 To design..., To develop..., To evaluate..."),
    ("P15", "effectiveness and relevance of the generated answers", "effectiveness 含义模糊。", "accuracy, relevance, and contextual consistency of generated answers"),
    ("P15", "assess accuracy and usability of the system", "缺冠词。", "assess the accuracy and usability of the system"),
    ("P16", "issues regarding faculty members and students", "范围表达不清。", "academic and student-service matters"),
    ("P16", "especially to international students", "介词错误。", "especially for international students"),
    ("P16", "With this respect", "固定搭配错误。", "In this respect"),
    ("P16", "the use of prompt containing system prompt", "缺冠词/复数且重复 prompt。", "the use of a prompt containing system instructions"),
    ("P16", "allows increasing answer accuracy", "搭配错误。", "helps increase answer accuracy"),
    ("P16", "applied in designing question answering system for university information service", "缺冠词且复合定语需连字符。", "applied to designing a question-answering system for university information services"),
    ("P16", "Through transformation of documents", "缺冠词。", "Through the transformation of documents"),
    ("P17", "query response system / query-response", "术语不自然。", "question-answering system"),
    ("P17", "the university and its clients", "clients 不符合学生服务语境。", "the university and its users/students"),
    ("P17", "subjects – undergraduate administration ... visa guide", "破折号和词语选择不当。", "areas: undergraduate administration, course scheduling, university services, visa guidance, and student life"),
    ("P17", "one needs to mention the following issues", "主观且冗长。", "The study covers the following components"),
    ("P17", "without any additions to it including open-ended dialogues outside of standard knowledge area", "语法和逻辑不清。", "and will exclude open-ended dialogue outside the defined knowledge domains"),
    ("P17", "RAG system including ... is utilized in this work", "缺冠词，主语过长，被动生硬。", "This study uses a RAG system that combines vector search with large-language-model response generation."),
    ("P17", "large scale deployment", "复合形容词需连字符。", "large-scale deployment"),
    ("P17", "This thesis paper consists of five chapters", "与当前目录只有三章不一致。", "按 P1 提交要求改写，或补齐/更新目录与后续章节。"),
]


CHAPTER_2 = [
    ("P18", "useful in helping with accessing information", "搭配生硬。", "useful in helping users access information"),
    ("P18", "answering questions of the users", "表达不自然。", "answering user questions"),
    ("P18", "The study by Saha & Saha (2024)", "叙述性引用不应使用 &（取决于格式，但全文也不统一）。", "Saha and Saha (2024)"),
    ("P19", "Malaysian higher educational institutions", "固定搭配错误。", "Malaysian higher education institutions"),
    ("P19", "availability ... requires knowledge of more than one language and platform", "逻辑不清：availability 不会要求知识。", "accessing this information often requires navigating multiple languages and platforms"),
    ("P19", "impact ... on the work of students", "work 指代不合适。", "impact ... on students' academic adjustment and daily activities"),
    ("P19", "GPT-4 can provide impressive conversational skills: understand...", "冒号后结构不平行。", "GPT-4 demonstrates strong conversational capabilities: it can understand..."),
    ("P19", "the wide range of LLM's abilities", "所有格/单复数错误。", "the wide range of LLM capabilities / LLMs' capabilities"),
    ("P19", "Researches claim", "research 不可数；researches 不适用于此处。", "Research indicates / Studies suggest"),
    ("P20", "A retrieval of relevant data is performed by the RAGs", "不自然且 RAGs 用法错误。", "RAG systems retrieve relevant data from an updatable external database"),
    ("P20", "models. (Lewis et al., 2020).", "引用前多余句号；此问题在 P20 多次出现。", "models (Lewis et al., 2020)."),
    ("P20", "accuracy of model's question-answering performance", "缺冠词且语义重复。", "the model's question-answering accuracy"),
    ("P20", "loading it as cue words for the model", "cue words 不准确。", "providing the retrieved content as context to the model"),
    ("P20", "It allows enhancing ... accuracy and professionalism", "allow 用法错误；professionalism 不适合回答质量。", "This improves the accuracy and domain specificity of the answers"),
    ("P20", "two common types of retrieval such as", "既说 two types 又用 such as，结构冲突。", "two common retrieval methods: BM25 sparse retrieval and dense retrieval"),
    ("P20", "based on semantic similarities", "通常用不可数。", "based on semantic similarity"),
    ("P20", "facilitate the process of generation of an answer", "名词化过重。", "facilitate answer generation"),
    ("P20", "RAG can ... improve ... but it depends", "it 指代不清。", "RAG can improve...; however, its performance depends on retrieval and context selection."),
    ("P20", "search materials", "术语不专业。", "retrieved evidence / retrieved documents"),
    ("P21", "Intelligent Question Answering Systems in Education", "该二级标题大小写与其他全大写标题不一致。", "按统一标题样式改为全大写或统一为 Title Case。"),
    ("P21", "provide help with learning, feedback, administrative queries, and provision of information", "并列结构不平行。", "support learning, provide feedback, answer administrative queries, and deliver academic information"),
    ("P21", "in the case of narrow pre-defined domain", "缺冠词且表达生硬。", "within a narrowly defined domain"),
    ("P21", "ambiguity of intentions", "术语不自然。", "ambiguous user intent"),
    ("P21", "There were several studies that indicated", "时态/结构冗长。", "Several studies have indicated"),
    ("P21", "presentation of the answer by Unimib Assistant", "搭配不自然。", "presentation of answers provided by the Unimib Assistant"),
    ("P21", "unworkable links", "词语不准确。", "broken or non-functional links"),
    ("P21", "As seen, these results", "口语化。", "These results"),
    ("P22", "most of reported university chatbots", "缺冠词。", "most reported university chatbots"),
    ("P22", "resources in another one", "one 多余。", "resources in another language"),
    ("P22", "a conversational artificial intelligence", "AI 在此不可数。", "conversational artificial intelligence / a conversational AI system"),
    ("P22", "among the criteria ... is the possibility", "主谓/逻辑生硬。", "one accessibility criterion is the ability"),
    ("P22", "It is possible due to", "It 指代不清。", "This capability is enabled by"),
    ("P22", "questions formulated by Malaysian students", "与研究对象“中国国际学生”不一致。", "questions submitted by students in Malaysia / Chinese international students"),
    ("P22", "enable comparing documents", "搭配错误。", "enable comparison of documents"),
    ("P22", "whether they were created in English", "文本通常写成某种语言。", "whether they were written in English"),
    ("P23", "The recent research demonstrated", "冠词不自然。", "Recent research has demonstrated"),
    ("P23", "improving cross-lingual search for information retrieval accuracy and factuality", "修饰关系混乱。", "improving the accuracy and factual grounding of cross-lingual information retrieval"),
    ("P23", "the document written in Malay or English", "应为复数。", "documents written in Malay or English"),
    ("P23", "Qwen3-Max native", "词序错误。", "native Qwen3-Max"),
    ("P23", "knowledge bases (Gao et al., 2023)", "段末缺句号。", "knowledge bases (Gao et al., 2023)."),
    ("P24", "RAG-based Systems", "标题大小写不统一。", "RAG-Based Systems"),
    ("P24", "Very high / Lower", "表内比较结论缺少判定依据或来源说明。", "增加 “Source: Author's comparison based on...” 或明确评价标准。"),
    ("P25", "international Chinese students' studying at the Malaysian universities", "所有格和冠词错误。", "Chinese international students studying at Malaysian universities"),
    ("P25", "too few works", "论文语境建议用 studies。", "too few studies"),
    ("P25", "Secondly", "与 First 不平行。", "Second"),
    ("P25", "successful operations in trilingual interaction", "搭配不自然。", "successful operation in trilingual settings"),
    ("P25", "there exist significant gaps", "冗长。", "significant gaps remain"),
    ("P25", "heterogeneous sources' integration", "所有格表达生硬。", "the integration of heterogeneous sources"),
    ("P25", "rely on the external retrieval", "冠词多余。", "rely on external retrieval"),
]


CHAPTER_3 = [
    ("P27", "there arises a need for the development", "冗长。", "there is a growing need to develop"),
    ("P27", "cannot cater to such information needs of students that require", "关系从句指代不清。", "cannot meet students' need for immediate information about"),
    ("P27", "Retrieval Augmented Generation", "缺连字符。", "Retrieval-Augmented Generation"),
    ("P27", "RAGs retrieve", "RAGs 用法不统一。", "RAG systems retrieve"),
    ("P27", "ensures contextual grounding and prevents hallucinations", "绝对化。", "improves contextual grounding and reduces the risk of hallucination"),
    ("P27", "Studies in recent times have proven", "口语且 proven 过强。", "Recent studies have demonstrated"),
    ("P28", "a specialized RAG designed specifically for serving FTSM, UKM, the Faculty...", "重复且机构关系不清。", "a specialized RAG system designed for FTSM at UKM"),
    ("P28", "multilinguism", "拼写错误。", "multilingualism"),
    ("P28", "applications of RAGs on campuses", "RAGs 用法不统一。", "campus RAG applications / RAG systems on campuses"),
    ("P28", "The rest of the chapter has been structured", "时态冗长。", "The remainder of the chapter is structured"),
    ("P28", "text embedding using DashScope model", "缺冠词。", "text embedding using the DashScope model"),
    ("P28", "doing Traditional to Simplified Chinese conversion", "口语化。", "converting Traditional Chinese to Simplified Chinese"),
    ("P28", "synonyms substitution", "单复数错误。", "synonym substitution"),
    ("P28", "Section 3.6 talks about", "偏口语。", "Section 3.6 describes"),
    ("P28", "The first boundary represents the users' interaction ... forwarding ... taking care...", "句子过长，分词结构主语不清。", "拆成三句，分别说明 UI、Application Service 和 Conversation Storage 的职责。"),
    ("P28", "requests' validation", "不自然的所有格。", "request validation"),
    ("P28", "stored to the Conversation Storage", "介词错误。", "stored in conversation storage"),
    ("P29", "uses vector database and knowledge base", "缺冠词。", "uses the vector database and knowledge base"),
    ("P29", "knowledge base updates are stored to the Conversation Storage", "介词错误，且知识库更新是否应存入会话存储需核实。", "stored in the appropriate data store；核对架构逻辑。"),
    ("P29", "In reality, the Application Service is implemented..., Langchain..., Vector Database..., and LLM...", "逗号拼接多个独立句，且大小写不统一。", "拆句；改为 In the implementation...；LangChain；vector database；the LLM。"),
    ("P29", "Tongyi Qwen retrieved via DashScope API", "retrieved 用词错误。", "Qwen3-Max accessed through the DashScope API"),
    ("P29", "described in sections 3.4 and 3.6 accordingly", "引用具体章节应大写；accordingly 用法不自然。", "described in Sections 3.4 and 3.6, respectively"),
    ("P29", "Figure 3.1 caption", "图题使用普通正文样式，和其他图题的 Caption 样式不一致。", "统一使用 Caption 样式并保持图题编号字段。"),
    ("P30", "five consecutive sub-stages ... will be saved", "substages 拼写形式和时态需统一。", "five consecutive substages ... are stored（或 will be stored，按项目状态统一）"),
    ("P30", "The input format ... is one of four – TXT, PDF, PNG, and JPG", "表达生硬，破折号用法不当。", "The pipeline supports four input formats: TXT, PDF, PNG, and JPG."),
    ("P30", "to see whether a SourceDocument metadata object needs to be created", "口语且逻辑不清。", "to determine whether a SourceDocument metadata object should be created"),
    ("P30", "parameters like", "正式写作建议。", "fields such as"),
    ("P30", "hashing is a suggested best practice", "suggested 语气和因果关系弱。", "hash-based duplicate detection is a common practice"),
    ("P30", "The obtained texts pass the filter", "搭配错误。", "The extracted text is passed to"),
    ("P30", "does not exceed configured chunk_size with an overlapping of configured chunk_overlap values", "语法错误。", "does not exceed the configured chunk_size and overlaps adjacent chunks by the configured chunk_overlap"),
    ("P30", "starts cutting large chunks of text", "偏口语。", "first splits text using larger structural separators"),
    ("P30", "only if it is required", "冗长。", "only when necessary"),
    ("P30", "diverse corporate documents", "研究对象是高校文件，corporate 不合适。", "diverse institutional documents"),
    ("P31", "using DashScope text-embedding-v4 model", "缺冠词。", "using the DashScope text-embedding-v4 model"),
    ("P31", "top retrieval efficiency in benchmark tests", "表达不自然且结论过强。", "strong retrieval performance in benchmark evaluations"),
    ("P31", "possibility to tune task instructions", "搭配错误。", "ability to configure task instructions"),
    ("P31", "the most versatile model of all in DashScope API library", "口语化且绝对化。", "a suitable model available through the DashScope API"),
    ("P31", "Obtained embeddings", "缺冠词。", "The resulting embeddings"),
    ("P32", "assistant for students at a multilingual institution, like UKM", "介词和举例表达不自然。", "assistant for students in a multilingual institution such as UKM"),
    ("P32", "generate inconsistently embedded vectors with poor alignment", "搭配不自然。", "produce embeddings that are poorly aligned with the stored chunk vectors"),
    ("P32", "extra whitespaces", "whitespace 通常不可数。", "extra whitespace"),
    ("P32", "converts Traditional Chinese into Simplified Chinese", "更自然的表达。", "converts text from Traditional Chinese to Simplified Chinese"),
    ("P32", "render semantically similar queries to have low alignment", "语法错误。", "cause semantically similar queries to have low embedding alignment"),
    ("P32", "Firstly / Secondly", "与其他章节 First/Second 不统一。", "First / Second"),
    ("P32", "supervisor can be replaced with penyelia or 导师", "扩展并非一定是替换。", "the query can be expanded to include penyelia and 导师"),
    ("P32", "knowledge domains like visas and permits", "like 偏口语。", "knowledge domains such as visas and permits"),
    ("P32", "research suggesting expanding a user's query", "结构生硬。", "research recommending that user queries be expanded"),
    ("P33-P34", "Figure 3.3", "3.5 节的 Agent 与 hybrid retrieval 工作流并不在 Figure 3.3。", "分别改为 Figure 3.4(a) 和 Figure 3.4(b)。"),
    ("P33", "erroneous answers about programmatic issues can negatively mislead the users", "programmatic 含义错误；negatively mislead 重复。", "incorrect answers about programme-related matters can mislead users"),
    ("P33", "Agent Initialization and Tool Decision.", "以正文粗体句开头，未使用正式子标题样式。", "若为 3.5 的子节，设为 3.5.1 并使用统一 Heading 样式；否则去掉标题式句号。"),
    ("P33", "makes the agent work under the ReAct approach", "搭配不自然。", "enables the agent to operate according to the ReAct framework"),
    ("P33", "conducts a ReAct-based tool decision / in case there is a need", "表达生硬。", "decides whether to invoke a tool using the ReAct framework / when domain knowledge is required"),
    ("P36", "Both the simulated context and query", "simulated 很可能为误词。", "Both the retrieved context and the query"),
    ("P36", "institutionalized response without any hallucinations", "institutionalized 用词错误且绝对化。", "institution-grounded response with a reduced risk of hallucination"),
    ("P36", "a Source Summary document is attached that consists of", "结构生硬。", "a source summary accompanies the response and lists"),
    ("P36", "students can validate their answers", "their answers 指代错误。", "students can verify the generated answers"),
    ("P36", "FastAPI Backend / Frontend Display / Conversation History Storage", "普通组件名大小写过度且不统一。", "FastAPI backend / frontend / conversation history storage，除非图中定义为正式组件名。"),
    ("P37", "\"What do I need to apply for the Master's program?\".", "问号后又加句号。", "\"What do I need to apply for the Master's programme?\""),
    ("P37", "extra embeddings generation", "名词搭配错误。", "additional embedding generation"),
    ("P37", "uses the Semantic Cache component", "普通组件名不必大写。", "uses a semantic-cache component"),
    ("P37", "previous solutions match ... via the surface form", "previous solutions 指代不清；介词不自然。", "conventional caches match queries based on surface forms"),
    ("P37", "The algorithm of the cache component", "搭配不自然。", "The cache algorithm"),
    ("P37", "generates an embedding vector out of it with the use of", "冗长。", "generates a query embedding using"),
    ("P37", "cosine similarity ... ranging from 0 to 1 ... same content", "技术表述过度简化；cosine similarity 一般可为 -1 到 1，具体范围取决于模型/归一化。", "说明本系统实际返回的范围；将 same content 改为 similar semantic meaning。"),
    ("P37", "finds the most similar pair of vectors", "系统是查询向量与缓存向量比较，不是任意向量对。", "identifies the cached vector most similar to the query vector"),
    ("P37", "it is referred to as cache hit / cache miss occurs", "缺冠词。", "this is referred to as a cache hit / a cache miss occurs"),
    ("P37", "get added ... minimize chances ... make hit rate significantly low", "口语/冠词/搭配问题。", "are added ... minimize the risk ... significantly reduce the hit rate"),
    ("P38", "assigned with Time-To-Live expiration period", "介词和大小写不统一。", "assigned a time-to-live (TTL) period"),
    ("P38", "do not become outdated in time", "语义不通。", "are not served after they become outdated"),
    ("P38", "maximum cache entry size constraint ... limit was exceeded ... items got deleted", "实际描述的是条目数量，不是单条大小；时态也不统一。", "a maximum number of cache entries; if the limit is exceeded, the least-recently-used entries are removed"),
    ("P38", "These restrictions guarantee", "guarantee 过强。", "These constraints help maintain"),
    ("P38", "The effects ... are obvious", "主观、口语。", "The component reduces"),
    ("P38", "amount of API calls", "可数名词。", "number of API calls"),
    ("P38", "/api/cache/stats endpoint is provided", "缺冠词。", "the /api/cache/stats endpoint provides"),
    ("P38", "The entrance point", "固定表达错误。", "The entry point"),
    ("P38", "opens up the locally hosted web service", "口语。", "starts/exposes the locally hosted web service"),
    ("P39", "capability of handling ... allows serving", "搭配生硬。", "ability to handle ... allows the application to serve"),
    ("P39", "via POST /api/chat endpoint", "缺冠词。", "via the POST /api/chat endpoint"),
    ("P39", "responds with the StreamingResponse providing", "结构生硬。", "returns a StreamingResponse that delivers"),
    ("P39", "SSE is a standard of the WHATWG", "搭配错误。", "SSE is defined in the WHATWG HTML Living Standard"),
    ("P39", "through the unidirectional HTTP connection", "冠词不当。", "through a unidirectional HTTP connection"),
    ("P39", "handles the chat message: invokes ... as a part of one single", "冒号后缺明确主语；one single 重复。", "handles each chat message by invoking... as part of a single request-processing flow"),
    ("P39", "Agent/ holds", "实际项目目录为小写 agent/，且全文路径大小写应严格一致。", "agent/ contains"),
    ("P40", "Integration with DashScope API", "缺冠词。", "Integration with the DashScope API"),
    ("P40", "Embedding model is fixed", "缺冠词。", "The embedding model is fixed"),
    ("P40", "models names", "名词形式错误。", "model names"),
    ("P40", "through CHAT_MODEL_NAME environment variable", "缺冠词。", "through the CHAT_MODEL_NAME environment variable"),
    ("P40", "via Settings Page", "普通页面名大小写不一致。", "via the Settings page"),
    ("P40", "stored in data/... directory", "多处目录名前缺 the；id 建议写 ID。", "stored in the data/... directory；ID"),
    ("P40", "created_at time / updated_at time", "time 重复。", "created_at timestamp / updated_at timestamp"),
    ("P40", "each entry ... having a similarity threshold ... TTL ... limit", "悬垂修饰；这些是缓存配置，不是每条记录的字段。", "The cache is configured with a similarity threshold of 0.92, a seven-day TTL, and a limit of 500 entries."),
    ("P40", "with collection named ftsm_rag_agent", "缺冠词。", "with a collection named ftsm_rag_agent"),
    ("P40", "was consciously chosen", "不自然。", "was deliberately chosen"),
    ("P40", "complying with the privacy requirements", "绝对化且缺少具体标准。", "which helps the system meet local data-handling requirements"),
    ("P40", "knowledge base trained anew", "搭配错误。", "the knowledge base is retrained"),
    ("P41", "provides / started / was outlined / describes", "同一总结段时态跳变。", "统一使用过去时总结本章，或统一使用现在时描述论文结构。"),
    ("P41", "the vector database, knowledge base, and conversation storage", "并列冠词不统一。", "the vector database, the knowledge base, and conversation storage"),
    ("P41", "semantic diversification by expanding the query to its synonyms", "搭配不自然。", "semantic expansion using domain-specific synonyms"),
    ("P41", "factually correct and institutionally verified answers", "缺少实验章节支撑的绝对结论。", "better-grounded answers that can be checked against institutional sources"),
    ("P41", "the stored answers returned", "缺谓语。", "the stored answers are returned"),
    ("P41", "the necessity to ensure..., multilingual ... capabilities, as well as efficient performance", "并列结构不平行。", "ensuring factual grounding, supporting multilingual queries, and maintaining efficiency"),
]


REFERENCES = [
    ("P42-P45", "All entries use Normal style", "没有悬挂缩进，连续条目视觉密度高。", "按 UKM 模板应用统一 Reference style 和 hanging indent。"),
    ("P42-P45", "pp. 1877-1901 / pp. 758–759 / pp. 333–389", "页码范围连接号混用。", "统一使用学院要求的连字符或 en dash。"),
    ("P42-P45", "Body citations use both & and and", "正文与参考文献作者连接方式不统一。", "按一种引用体系统一正文与文末格式。"),
    ("P42", "Faculty of Information Science and Technology (FTSM), UKM (2024)", "该条目在正文中未找到对应引用。", "若确实未使用则删除；若用于机构背景或数据来源，在相应正文处补引。"),
    ("P42-P45", "Title capitalization", "论文题名、网页题名和期刊信息的大小写规则看起来不完全统一。", "按 UKM 指定 Harvard 样式统一 sentence case/title case；原始专名保持不变。"),
    ("P42-P45", "URLs and accessed dates", "格式总体完整，但网页、arXiv、DOI 的呈现方式不完全一致。", "统一 DOI/URL 顺序、Available at 和 Accessed 日期格式。"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_font(run, size=10.5, bold=False, color=TEXT, name="Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, size=10.5, bold=False, color=TEXT) -> None:
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)


def set_paragraph_format(paragraph, after=4, before=0, line=1.15) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "论文审阅问题清单  |  ", size=8.5, color=MUTED)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, size=8.5, color=MUTED)


def add_header(section) -> None:
    paragraph = section.header.paragraphs[0]
    add_text(
        paragraph,
        "Hou Zhengbo P1 | Language, Expression and Formatting Review",
        size=8.5,
        color=MUTED,
    )


def add_issue_table(doc: Document, title: str, rows, start_no: int) -> int:
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.45, 0.65, 2.45, 2.0, 4.25]
    headers = ["No.", "页", "原文/位置", "问题", "建议修改"]
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(p, after=0, line=1.0)
        add_text(p, header, size=8.5, bold=True, color=TEXT)
    set_repeat_table_header(table.rows[0])

    number = start_no
    for page, original, problem, suggestion in rows:
        row = table.add_row()
        set_row_cant_split(row)
        values = [str(number), page, original, problem, suggestion]
        for index, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if number % 2 == 0:
                set_cell_shading(cell, "FAFBFC")
            p = cell.paragraphs[0]
            set_paragraph_format(p, after=0, line=1.05)
            if index in (0, 1):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(
                p,
                value,
                size=8.2,
                bold=index == 0,
                color=RED if index == 3 and "错误" in problem else TEXT,
            )
        number += 1
    doc.add_paragraph()
    return number


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    configure_styles(doc)
    add_header(section)
    add_footer(section)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(title, after=4)
    add_text(title, "论文语言与格式审阅问题清单", size=22, bold=True, color=TEXT)

    subtitle = doc.add_paragraph()
    set_paragraph_format(subtitle, after=12)
    add_text(
        subtitle,
        "HouZhengbo P1 降低版.docx | 审阅范围：拼写、语法、措辞、前后一致性、目录、图表、引用与版式",
        size=11,
        color=MUTED,
    )

    note = doc.add_paragraph()
    set_paragraph_format(note, after=8, line=1.2)
    add_text(note, "审阅口径：", bold=True)
    add_text(
        note,
        "按授课型硕士论文标准处理，不以期刊审稿的严苛程度评价创新性；重点是让文本清楚、自然、统一、可提交。页码为 Word/PDF 物理页码，原文短语可直接在文档中搜索。",
    )

    summary = doc.add_paragraph()
    set_paragraph_format(summary, after=12, line=1.2)
    add_text(summary, "总体判断：", bold=True)
    add_text(
        summary,
        "结构和排版整体可读，45 页中未发现文字重叠、图片截断或表格溢出。主要问题集中在中式英语、冠词/搭配、方法章节时态、交叉引用和目录字段。先修正下列“必须先改”项目，再按章节清单逐项处理即可。",
    )

    doc.add_heading("必须先改", level=1)
    number = add_issue_table(doc, "高优先级问题", CRITICAL, 1)

    doc.add_heading("全文统一", level=1)
    number = add_issue_table(doc, "重复出现的语言与格式问题", GLOBAL, number)

    doc.add_heading("逐页清单", level=1)
    number = add_issue_table(doc, "前置部分与摘要", FRONT, number)
    number = add_issue_table(doc, "Chapter I: Introduction", CHAPTER_1, number)
    number = add_issue_table(doc, "Chapter II: Literature Review", CHAPTER_2, number)
    number = add_issue_table(doc, "Chapter III: Methodology", CHAPTER_3, number)
    number = add_issue_table(doc, "References", REFERENCES, number)

    doc.add_heading("处理顺序建议", level=1)
    for text in (
        "更新目录、图表目录和交叉引用，删除 P26 的斜杠与 P10 的空 Pictures 栏目。",
        "确定项目状态：已完成就统一过去/现在时；P1 方案就统一将来时，并处理 Chapter 4/5 的说明。",
        "统一引用体系、术语大小写、缩略语和 programme/program 的英式英语规则。",
        "按逐页清单修句，最后在 Word 中执行 Ctrl+A 后按 F9 更新全部字段。",
        "重新检查目录页码、图表编号和参考文献悬挂缩进。",
    ):
        p = doc.add_paragraph(style="List Number")
        set_paragraph_format(p, after=4, line=1.15)
        add_text(p, text)

    ending = doc.add_paragraph()
    set_paragraph_format(ending, before=8, after=0)
    add_text(
        ending,
        f"本报告共列出 {number - 1} 个具体或全局问题。专有名词、作者姓名和代码标识未按普通拼写错误处理。",
        size=9.5,
        color=MUTED,
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
